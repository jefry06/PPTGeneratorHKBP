from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import re
import urllib.parse

from GetFromInternet import get_lyrics_from_url, get_pericope_from_url



class SundayService:
    def __init__(self, title=None, datetime=None, topic=None, sequences=None):
        self.title = title
        self.datetime = datetime
        self.topic = topic
        self.sequences = sequences or []

def add_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    for shape in slide.shapes:
        if shape.is_placeholder:
            sp = shape.element
            sp.getparent().remove(sp)

    current_dir = Path(__file__).parent
    image_header = current_dir.parent / "resources" / "header.png"
    slide.shapes.add_picture(str(image_header), Inches(0.2), Inches(0.2))

    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5))
    text_frame = text_box.text_frame
    text_frame.word_wrap = True
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    return text_frame

def add_paragraphs_to_slide(prs, paragraphs, font_size_map):
    for paragraph in paragraphs:
        text_frame = add_slide(prs)
        p = text_frame.paragraphs[0]
        p.text = paragraph
        p.alignment = PP_ALIGN.CENTER
        paragraph_length = len(paragraph)

        for length, size in font_size_map.items():
            if paragraph_length > length:
                p.font.size = Pt(size)
                break
        else:
            p.font.size = Pt(font_size_map[0])

def read_file(file_path):
    # Open the file
    with open(file_path, 'r') as file:
        # Read lines from the file
        lines = file.readlines()

    # Initialize variables
    title = lines[0].strip().replace(",", "\n")  # First line (Title)
    datetime = lines[1].strip()  # Second line (Datetime)
    topic = lines[2].split(';')[1].strip()  # Topic after splitting on ';'

    # Initialize a list for other items
    worship_service = []
    worship_service.append(["Opening Song", clean_text(lines[3].strip())])
    worship_service.append(["Votum", clean_text(lines[4].strip())])
    worship_service.append(["Opening Song 2", clean_text(lines[5].strip())])
    worship_service.append(["Law of God", clean_text(lines[6].strip())])
    worship_service.append(["Praise Song 1", clean_text(lines[7].strip())])
    worship_service.append(["Confession of Sins", clean_text(lines[8].strip())])
    worship_service.append(["Praise Song 2", clean_text(lines[9].strip())])
    worship_service.append(["Epistle Reading", clean_text(lines[10].strip())])
    worship_service.append(["Praise Song 3", clean_text(lines[11].strip())])
    worship_service.append(["Creed", clean_text(lines[12].strip())])
    worship_service.append(["Announcements", clean_text(lines[13].strip())])
    worship_service.append(["Offering Song", clean_text(lines[14].strip())])
    worship_service.append(["Sermon", clean_text(lines[15].strip())])
    worship_service.append(["Closing Song", clean_text(lines[16].strip())])
    worship_service.append(["Closing Prayer", clean_text(lines[17].strip())])

    # for item in worship_service:
    #     print(item)

    sunday_service = SundayService(title, datetime, topic, worship_service);
    return sunday_service


def clean_text(text):
    text = text.replace('â€“', '-')
    text = text.replace('â€¦', '...')
    cleaned_text = re.sub(r'^\d+\.\s*\t?', '', text)
    return cleaned_text


def create_caption(prs, text):
    print(f"Generate Slide : {text.replace("\n", " ")}")
    paragraphs = []
    paragraphs.append(text.upper())

    font_size_map = {250: 36, 200: 40, 50: 48, 0: 54}
    add_paragraphs_to_slide(prs, paragraphs, font_size_map)


def create_song_caption(prs, text, opening_song=False, language="1"):
    if language == "1":
        text = text.replace("BE.", "MARENDE\nBUKU ENDE ")
    else:
        text = text.replace("NH.", "BERNYANYI\nNYANYIAN HKBP ")
    create_caption(prs, text)
    generate_song_lyric(prs, text, opening_song, language)


def generate_song_lyric(prs, text, opening_song, language):
    paragraphs = []
    song_type = "BUKU ENDE" if language == "1" else "NYANYIAN HKBP"

    song_number, start_verse, sign_verse, last_verse = 1, 1, None, 1
    pattern = r"(BUKU ENDE|NYANYIAN HKBP)\s+(\d+)\s*:\s*(\d+)\s*([-+–…])\s*(\d+|…)"
    pattern_full_song = r"(BUKU ENDE|NYANYIAN HKBP)\s*(\d+)"
    match = re.search(pattern, text)

    if match:
        song_type = match.group(1)  # Either "BUKU ENDE" or "NYANYIAN HKBP"
        song_number = int(match.group(2))  # Song number
        start_verse = int(match.group(3))  # Start verse
        sign_verse = match.group(4)  # Sign verse
        last_verse = int(match.group(5))   # Last verse
    else:
        match = re.search(pattern_full_song, text)
        if match:
            song_number = int(match.group(2))
            start_verse = 1
            last_verse = 100

    lyrics = get_lyrics_from_url(song_type, song_number, start_verse, last_verse)
    counter = 0

    # Sort lyrics by verse number and filter verses within range
    sorted_lyrics = sorted(lyrics, key=lambda x: x[0])
    max_verse = sorted_lyrics[-1][0] if sorted_lyrics else 0
    last_verse = min(last_verse, max_verse)

    for verse_number, verse_text in sorted_lyrics:
        if sign_verse == "+":
            if verse_number == start_verse or verse_number == last_verse:
                if counter % 2 == 0 and counter != 0 and not opening_song:
                    paragraphs.append("MUSIK")
                counter += 1
                paragraphs.append(verse_text)
        else:
            if start_verse <= verse_number <= last_verse:
                if counter % 2 == 0 and counter != 0 and not opening_song:
                    paragraphs.append("MUSIK")
                counter += 1
                paragraphs.append(verse_text)

    font_size_map = {250: 36, 200: 40, 50: 48, 0: 54}
    add_paragraphs_to_slide(prs, paragraphs, font_size_map)
def create_sins_caption(prs, text):
    if "M" in text:
        text = "MANOPOTI DOSA"
    else:
        text = "PENGAKUAN DOSA"
    create_caption(prs, text)

def create_epistel_caption(prs, text):
    pericope = text
    text = re.sub(r'(?i)Epistel\s*:', 'EPISTEL\n', text)
    create_caption(prs, text)
    generate_pericope(prs, pericope)

def create_creed_caption(prs, text):
    if "haporseaon" in text.lower():
        text = "MANGHATINDANGHON \nHAPORSEAON"
    else:
        text = "PENGAKUAN IMAN"
    create_caption(prs, text)

def create_sermon_caption(prs, text):
    pericope = text
    text = re.sub(r'(?i)Jamita\s*:', 'JAMITA\n', text)
    text = re.sub(r'(?i)Kotbah\s*:', 'KHOTBAH\n', text)
    text = re.sub(r'(?i)Khotbah\s*:', 'KHOTBAH\n', text)
    create_caption(prs, text)
    generate_pericope(prs, pericope)

def generate_pericope(prs, pericope):
    pericope_path, start_verse, end_verse = extract_pericope(pericope)

    paragraphs = get_pericope_from_url(pericope_path, start_verse, end_verse)
    font_size_map = {250: 36, 200: 40, 50: 48, 0: 54}
    add_paragraphs_to_slide(prs, paragraphs, font_size_map)

def extract_pericope(text):
    pattern = r"([A-Za-z\s]+)\s(\d+)\s:\s(\d+)\s-\s(\d+)"
    # Search for the pattern in the line
    match = re.search(pattern, text)

    if match:
        book = match.group(1).strip().upper()  # Convert book name to uppercase
        chapter = match.group(2)
        start_verse = match.group(3)
        end_verse = match.group(4)

        # Format the pericope like "MATIUS 24:9-14"
        pericope = f"{book} {chapter}:{start_verse}-{end_verse}"

        # Construct the URL encoding format like "MATIUS/24/9/14"
        url_encoded_pericope = f"{book}/{chapter}/{start_verse}/{end_verse}"

        # URL encode the result
        encoded_pericope = urllib.parse.quote(url_encoded_pericope)
        return encoded_pericope, start_verse, end_verse
    else:
        return None, None

def generate_ppt():
    prs = Presentation()

    file_path = "f:/batak.txt"
    sunday_service = read_file(file_path)
    worship_service_dict = dict(sunday_service.sequences)

    create_caption(prs, sunday_service.title)
    create_song_caption(prs, worship_service_dict.get("Opening Song"), True)
    create_caption(prs, worship_service_dict.get("Votum"))
    create_song_caption(prs, worship_service_dict.get("Opening Song 2"))
    create_caption(prs, worship_service_dict.get("Law of God"))
    create_song_caption(prs, worship_service_dict.get("Praise Song 1"))
    create_sins_caption(prs, worship_service_dict.get("Confession of Sins"))
    create_song_caption(prs, worship_service_dict.get("Praise Song 2"))
    create_epistel_caption(prs, worship_service_dict.get("Epistle Reading"))
    create_song_caption(prs, worship_service_dict.get("Praise Song 3"))
    create_creed_caption(prs, worship_service_dict.get("Creed"))
    create_caption(prs, worship_service_dict.get("Announcements"))
    create_song_caption(prs, worship_service_dict.get("Offering Song"))
    create_sermon_caption(prs, worship_service_dict.get("Sermon"))
    create_song_caption(prs, worship_service_dict.get("Closing Song"))
    create_caption(prs, worship_service_dict.get("Closing Prayer"))

    prs.save("f:/newSlide.pptx")

# generate_ppt()